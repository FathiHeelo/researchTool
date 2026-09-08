"""Exports never rerun evaluation or alter evidence."""
import csv
import io
import json
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from app.evaluation.research import records, metadata, analytics


def scalar(value):
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    return value


def excel_scalar(value):
    value = scalar(value)
    if isinstance(value, str):
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', lambda m: '\\u%04x' % ord(m.group()), value)
        if len(value) > 32767:
            raise ValueError('A research field exceeds the Excel cell limit. Export CSV to preserve the full text.')
    return value


def detailed(rows):
    output = []
    for r, d in rows:
        case = r.snapshot.get('case', {})
        response = r.snapshot.get('response', {})
        item = {'Result ID': r.id, 'Run ID': r.run_id, 'Case ID': r.case_id, 'Requirement': r.requirement,
                'Expected Rule': case.get('expected_rule'), 'Model': r.model,
                'Source Row': case.get('original_row_number', case.get('source_row', case.get('source_row_index'))),
                'Model Source Column': response.get('source_column_name'), 'Generated Output': response.get('generated_output'),
                'Source Snapshot': r.snapshot}
        item['Evaluation Provenance'] = r.evaluation.get('provenance')
        item.update({f'Metadata: {k}': v for k,v in metadata(r).items()})
        for label, values in [('Automated', r.scores), ('Override', d['override_scores']), ('Accepted', d['accepted_scores'])]:
            item.update({f'{label}: {k}': v for k,v in values.items()})
        item.update({'Accepted Overall Accuracy': d['accepted_overall_accuracy'], 'Automated Overall Accuracy': r.overall_accuracy,
                     'Execution State': r.evaluation.get('execution', r.scores.get('execution')),
                     'Hallucination': r.hallucination_detected, 'Hallucinated Functions': r.hallucinated_functions,
                     'Automated Error Tags': r.error_tags, 'Accepted Error Tags': d['error_tags'], 'Notes': r.notes,
                     'Review Note': d['review_note'], 'Ground Truth Warning': d['ground_truth_warning'],
                     'Ground Truth Warning Count': d['ground_truth_warning_count']})
        output.append(item)
    return output


def safe_csv(value):
    value = scalar(value)
    # Spreadsheet clients must not execute imported research strings as formulas.
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def export_run(session, project, run, format, include_ground_truth_warnings=True, strategy_key=None, dimension_key=None, sanitizer=None):
    rows, _ = records(session, project.id, run.id, include_ground_truth_warnings=include_ground_truth_warnings)
    details = detailed(rows)
    if sanitizer: details = sanitizer(details)
    if format == 'csv':
        stream = io.StringIO(newline='')
        keys = list(dict.fromkeys(k for item in details for k in item)) or ['Result ID', 'Case ID', 'Model']
        writer = csv.DictWriter(stream, fieldnames=keys)
        writer.writeheader()
        writer.writerows({k: safe_csv(v) for k,v in item.items()} for item in details)
        return stream.getvalue().encode('utf-8-sig')
    analysis = analytics(session, project.id, run.id, include_ground_truth_warnings=include_ground_truth_warnings,
                         strategy_key=strategy_key, dimension_key=dimension_key)
    book = Workbook(); book.remove(book.active)
    def sheet(name, items):
        if sanitizer: items = sanitizer(items)
        ws = book.create_sheet(name)
        keys = list(dict.fromkeys(k for item in items for k in item)) or ['Information']
        ws.append(keys)
        for item in items:
            ws.append([excel_scalar(item.get(k)) for k in keys])
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            ws.column_dimensions[cell.column_letter].width = min(55, max(18, len(str(cell.value))+2))
        for index, key in enumerate(keys, 1):
            if any(term in key.lower() for term in ('accuracy', 'reliability', 'rate', 'percentage', 'delta')):
                for row in ws.iter_rows(min_row=2, min_col=index, max_col=index):
                    if isinstance(row[0].value, (int, float)): row[0].number_format = '0.00"%"'
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    sheet('Run Summary', [{'Project': project.name, 'Run ID': run.id, 'Timestamp': run.created_at.isoformat(),
          'Evaluator Version': run.summary.get('evaluator_version', 'legacy'),
          **{k:v for k,v in analysis.items() if not isinstance(v, (list, dict))},
          'Metric Configuration': run.summary.get('metric_configuration', 'Legacy equal weights over stored keys'),
          'Definitions': 'Accuracy: run-weighted normalized accepted scores. Reliability: accepted execution=1 / eligible responses. Hallucination: invented function responses / responses.'}])
    sheet('Detailed Results', details)
    sheet('Model Summary', analysis['models'])
    sheet('Metric Summary', [{'Metric': k, 'Accepted Average': v} for k,v in analysis['average_component_scores'].items()])
    if analysis['strategies']: sheet('Strategy Analysis', analysis['strategies'] + analysis['strategy_deltas'])
    if analysis['dimensions']: sheet('DQ Dimension Analysis', analysis['dimensions'])
    if analysis['agreement']: sheet('Agreement Matrix', analysis['agreement'])
    sheet('Error Analysis', analysis['errors'])
    sheet('Ground Truth Warnings', [{'Case': r.case_id, 'Model': r.model, 'Requirement': r.requirement,
           'Expected Rule': r.snapshot.get('case', {}).get('expected_rule'), **w} for r,d in rows for w in d['ground_truth_warnings']])
    audits = []
    for r,d in rows:
        for audit in d['audit_history']:
            for field in set(audit.previous_value) | set(audit.new_value):
                if audit.previous_value.get(field) != audit.new_value.get(field):
                    audits.append({'Result': r.id, 'Case': r.case_id, 'Model': r.model, 'Actor': audit.actor,
                          'Timestamp': audit.created_at.isoformat(), 'Field': field, 'Previous': audit.previous_value.get(field),
                          'New': audit.new_value.get(field), 'Reason': audit.reason})
    sheet('Manual Overrides', audits)
    sheet('Metric Configuration', run.summary.get('metric_configuration', [{'Policy': 'Legacy equal weights over stored keys'}]))
    stream = io.BytesIO(); book.save(stream); return stream.getvalue()
