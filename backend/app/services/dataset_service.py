import csv
import re
from datetime import date, datetime, time
import math
from io import TextIOWrapper
from pathlib import PurePath
from zipfile import ZipFile

from fastapi import UploadFile
from openpyxl import load_workbook
from starlette.concurrency import run_in_threadpool

from app.schemas.dataset import CSVUploadMetadata, SelectedXLSXUploadMetadata, XLSXUploadMetadata


class InvalidUpload(ValueError):
    def __init__(self, message: str, code: str = "unreadable_dataset", status_code: int = 400):
        super().__init__(message)
        self.code = code
        self.status_code = status_code


def validate_csv(stream):
    text = TextIOWrapper(stream, encoding="utf-8-sig", newline="")
    width = None
    try:
        for row in csv.reader(text, strict=True):
            if any(any(ord(char) < 32 and char not in "\t\r\n" for char in cell) for cell in row):
                raise InvalidUpload("CSV contains unreadable control characters", "malformed_csv")
            if not any(cell.strip() for cell in row):
                continue
            if width is None:
                width = len(row)
            elif len(row) != width:
                raise InvalidUpload("CSV rows must contain a consistent number of fields", "malformed_csv")
    except (UnicodeError, csv.Error) as error:
        raise InvalidUpload("CSV must contain readable UTF-8 text with valid CSV quoting", "malformed_csv") from error
    finally:
        text.detach()
    if width is None:
        raise InvalidUpload("CSV contains no usable tabular data", "unreadable_dataset")


def validate_xlsx(stream):
    workbook = None
    usable = False
    try:
        # Check archive integrity, including entries not consumed by openpyxl.
        with ZipFile(stream) as archive:
            if archive.testzip() is not None:
                raise ValueError("Invalid workbook archive")
        stream.seek(0)
        workbook = load_workbook(stream, read_only=True, data_only=False)
        # Include empty and hidden worksheets, in workbook order; exclude chartsheets.
        sheet_names = [worksheet.title for worksheet in workbook.worksheets]
        for worksheet in workbook.worksheets:
            # Do not trust dimensions declared by third-party workbook writers.
            worksheet.reset_dimensions()
            for row in worksheet.iter_rows(values_only=True):
                if any(value is not None and (not isinstance(value, str) or value.strip()) for value in row):
                    usable = True
    except Exception as error:
        # Workbook parsers raise several XML/ZIP/value errors for untrusted files.
        raise InvalidUpload("XLSX is malformed or cannot be opened as an Excel workbook", "malformed_xlsx") from error
    finally:
        if workbook is not None:
            workbook.close()
    if not usable:
        raise InvalidUpload("Excel workbook contains no usable worksheets", "unreadable_dataset")
    return sheet_names


def validate_content(stream, extension):
    try:
        stream.seek(0)
        if extension == ".csv":
            validate_csv(stream)
        else:
            return validate_xlsx(stream)
    except OSError as error:
        raise InvalidUpload("Dataset could not be read", "unreadable_dataset") from error


def infer_cell_type(value):
    if value == "":
        return "unknown"
    if value.lower() in {"true", "false"}:
        return "boolean"
    if re.fullmatch(r"[+-]?\d+", value):
        return "integer"
    if re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value):
        return "decimal"
    try:
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            date.fromisoformat(value)
            return "date"
        if re.match(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}", value):
            datetime.fromisoformat(value)
            return "datetime"
    except ValueError:
        pass
    return "text"


def merge_types(current, incoming):
    if incoming == "unknown" or current == incoming:
        return current
    if current == "unknown":
        return incoming
    if {current, incoming} <= {"integer", "decimal"}:
        return "decimal"
    if {current, incoming} <= {"date", "datetime"}:
        return "datetime"
    return "text"


def parse_csv_metadata(stream):
    """Infer across all rows, retaining only ten previews and original text.

    Empty cells become JSON null; literal NA/NaN/null strings stay unchanged.
    Blank physical lines are skipped, but delimiter-only data rows are counted.
    """
    stream.seek(0)
    text = TextIOWrapper(stream, encoding="utf-8-sig", newline="")
    try:
        reader = csv.reader(text, strict=True)
        columns = next(row for row in reader if any(cell.strip() for cell in row))
        if len(set(columns)) != len(columns):
            raise InvalidUpload("CSV column names must be unique to represent rows as objects", "malformed_csv")
        types = dict.fromkeys(columns, "unknown")
        preview = []
        count = 0
        for row in reader:
            if not row:
                continue
            if len(row) != len(columns):
                if not any(cell.strip() for cell in row):
                    continue
                raise InvalidUpload("CSV rows must contain a consistent number of fields", "malformed_csv")
            count += 1
            for column, value in zip(columns, row):
                types[column] = merge_types(types[column], infer_cell_type(value))
            if len(preview) < 10:
                preview.append({column: value if value != "" else None for column, value in zip(columns, row)})
        return dict(total_row_count=count, column_count=len(columns), columns=columns,
                    inferred_data_types=types, preview_rows=preview)
    except (UnicodeError, csv.Error) as error:
        raise InvalidUpload("CSV must contain readable UTF-8 text with valid CSV quoting", "malformed_csv") from error
    except OSError as error:
        raise InvalidUpload("Dataset could not be read", "unreadable_dataset") from error
    finally:
        text.detach()


def infer_xlsx_value_type(value):
    """Classify actual values; None means no evidence, not incompatible evidence.

    Midnight datetimes are date-only; strings remain text, including yes/no.
    Openpyxl does not evaluate formulas; formula source remains text.
    """
    if value is None or isinstance(value, str) and not value.strip():
        return None
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, datetime):
        return "date" if value.time() == time(0) else "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        if not math.isfinite(value):
            return "unknown"
        return "integer" if value.is_integer() else "decimal"
    if isinstance(value, str):
        return "text"
    return "unknown"


def resolve_xlsx_types(observed):
    """Numeric/date families widen; incompatible evidence stays unknown."""
    if not observed:
        return "unknown"
    if len(observed) == 1:
        return next(iter(observed))
    if observed <= {"integer", "decimal"}:
        return "decimal"
    if observed <= {"date", "datetime"}:
        return "datetime"
    return "unknown"


def serialize_xlsx_value(value):
    """Keep scalar values; ISO-format temporal values and null nonfinite numbers."""
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def populated_positions(row):
    return [i + 1 for i, value in enumerate(row)
            if value is not None and (not isinstance(value, str) or value.strip())]


def xlsx_header_row(rows):
    """Return the first likely tabular header row and its zero-based index.

    Most workbooks start at row 1. Some exported research sheets include a
    title/preamble before the real header; when a later early row has multiple
    populated cells and following data, use that as the table start.
    """
    for index, row in enumerate(rows[:25]):
        if len(populated_positions(row)) >= 2 and any(populated_positions(next_row) for next_row in rows[index + 1:]):
            return index, list(row)
    return (0, list(rows[0])) if rows else (0, [])


def parse_xlsx_structure(stream, sheet_name):
    """The first likely table row is the header; count subsequent data rows.

    Width extends to the last populated cell across header and data, ignoring
    formatting-only cells. Preserve duplicate headers and use null for blanks,
    including columns with data but no header. No Pandas-generated names.
    Keep scalar header values; serialize date/time headers to ISO strings.
    An empty selected sheet has zero rows/columns (workbook validation still
    requires some usable data somewhere in the workbook).
    """
    workbook = None
    try:
        stream.seek(0)
        workbook = load_workbook(stream, read_only=True, data_only=False)
        worksheet = workbook[sheet_name]
        worksheet.reset_dimensions()
        rows = list(worksheet.iter_rows(values_only=True))
        header_index, header = xlsx_header_row(rows)
        width = 0
        row_count = 0
        observed_types = []
        preview = []
        for index, row in enumerate(rows[header_index:], start=header_index):
            populated = [i + 1 for i, value in enumerate(row)
                         if value is not None and (not isinstance(value, str) or value.strip())]
            if populated:
                width = max(width, populated[-1])
                if index > header_index:
                    row_count += 1
                    if len(preview) < 10:
                        preview.append((index - header_index, row))
            if index > header_index:
                while len(observed_types) < width:
                    observed_types.append(set())
                for position, value in enumerate(row[:width]):
                    inferred = infer_xlsx_value_type(value)
                    if inferred is not None:
                        observed_types[position].add(inferred)
        columns = header[:width] + [None] * max(0, width - len(header))
        columns = [value.isoformat() if isinstance(value, (date, datetime, time))
                   else None if isinstance(value, float) and not math.isfinite(value)
                   else value for value in columns]
        inferred_types = [resolve_xlsx_types(observed_types[i]) if i < len(observed_types) else "unknown"
                          for i in range(width)]
        return dict(total_row_count=row_count, column_count=width, columns=columns,
                    inferred_data_types=inferred_types,
                    preview_rows=[{
                        "row_index": index,
                        "values": [serialize_xlsx_value(row[i]) if i < len(row) else None
                                   for i in range(width)],
                    } for index, row in preview])
    except Exception as error:
        raise InvalidUpload("Selected worksheet could not be read", "unreadable_dataset") from error
    finally:
        if workbook is not None:
            workbook.close()


async def inspect_upload(file: UploadFile, sheet_name: str | None = None) -> CSVUploadMetadata | XLSXUploadMetadata:
    """Verify readability and return CSV metadata or XLSX worksheet discovery.

    A nonblank row is usable, including a header-only or single-column file.
    CSV input is comma-delimited UTF-8 (with an optional BOM).
    """
    if not file.filename or not file.filename.strip():
        raise InvalidUpload("A file with a filename is required")

    extension = PurePath(file.filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise InvalidUpload("Only .csv and .xlsx files are supported", "unsupported_file")
    if extension == ".csv" and sheet_name is not None:
        raise InvalidUpload("Sheet selection is only supported for XLSX files", "sheet_selection_not_supported")

    size = 0
    while chunk := await file.read(1024 * 1024):
        size += len(chunk)
    if size == 0:
        raise InvalidUpload("Uploaded file must not be empty", "empty_file")

    sheet_names = await run_in_threadpool(validate_content, file.file, extension)

    if extension == ".csv":
        metadata = await run_in_threadpool(parse_csv_metadata, file.file)
        return CSVUploadMetadata(filename=file.filename, extension=extension, size=size, **metadata)

    if sheet_name is not None:
        if sheet_name not in sheet_names:
            raise InvalidUpload("Requested worksheet does not exist; sheet names must match exactly", "sheet_not_found")
        structure = await run_in_threadpool(parse_xlsx_structure, file.file, sheet_name)
        return SelectedXLSXUploadMetadata(filename=file.filename, extension=extension, size=size,
                                         sheet_names=sheet_names, sheet_count=len(sheet_names),
                                         selected_sheet=sheet_name, **structure)

    return XLSXUploadMetadata(filename=file.filename, extension=extension, size=size,
                              sheet_names=sheet_names, sheet_count=len(sheet_names))
