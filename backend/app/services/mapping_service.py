import csv
from io import TextIOWrapper
from openpyxl import load_workbook

from app.services.dataset_service import InvalidUpload, serialize_xlsx_value, xlsx_header_row


def normalize_dataset(stream, extension, sheet_name, columns, configuration):
    assignments = {item.index: item for item in configuration.assignments}
    if len(assignments) != len(configuration.assignments):
        raise InvalidUpload('A source column cannot have multiple roles', 'invalid_mapping')
    if any(index >= len(columns) for index in assignments):
        raise InvalidUpload('Mapping references an unknown column', 'invalid_mapping')
    for role in ('requirement', 'expected_rule'):
        if sum(item.role == role for item in assignments.values()) != 1:
            raise InvalidUpload(f'Exactly one {role} column is required', 'invalid_mapping')
    if sum(item.role == 'case_id' for item in assignments.values()) > 1:
        raise InvalidUpload('Only one Case ID column is allowed', 'invalid_mapping')
    models = [{'id': f'column_{i}', 'label': item.label if item.label else str(columns[i]) if columns[i] is not None else f'Column {i + 1}',
               'source_column_index': i, 'source_column_name': columns[i]}
              for i, item in assignments.items() if item.role == 'model']
    if not models:
        raise InvalidUpload('At least one model output column is required', 'invalid_mapping')
    stream.seek(0)
    workbook = None
    text = None
    cases, responses = [], []
    try:
        if extension == '.xlsx':
            workbook = load_workbook(stream, read_only=True, data_only=False)
            worksheet = workbook[sheet_name]
            worksheet.reset_dimensions()
            all_rows = list(worksheet.iter_rows(values_only=True))
            header_index, _ = xlsx_header_row(all_rows)
            rows = iter(all_rows[header_index + 1:])
        else:
            text = TextIOWrapper(stream, encoding='utf-8-sig', newline='')
            rows = csv.reader(text, strict=True)
            next(row for row in rows if any(value.strip() for value in row))
        previous_csv_line = rows.line_num if extension == '.csv' else 0
        start_row = header_index + 2 if extension == '.xlsx' else 2
        for row_number, row in enumerate(rows, start=start_row):
            if extension == '.csv':
                row_number = previous_csv_line + 1
                previous_csv_line = rows.line_num
            if not row or (extension == '.xlsx' and not any(value is not None and (not isinstance(value, str) or value.strip()) for value in row)):
                continue
            values = [serialize_xlsx_value(row[i]) if i < len(row) else None for i in range(len(columns))]
            primary = {item.role: values[i] for i, item in assignments.items() if item.role in {'case_id', 'requirement', 'expected_rule'}}
            case_key = f'row_{row_number}'
            metadata = [{'source_column_index': i, 'source_column_name': columns[i],
                         'label': (assignments[i].label or columns[i]) if i in assignments else columns[i], 'value': value}
                        for i, value in enumerate(values) if i not in assignments or assignments[i].role == 'metadata']
            cases.append({'id': case_key, 'case_id': primary.get('case_id') if primary.get('case_id') not in (None, '') else case_key,
                          'requirement': primary['requirement'], 'expected_rule': primary['expected_rule'],
                          'original_row_number': row_number, 'metadata': metadata})
            for model in models:
                responses.append({'case_reference': case_key, 'model': model['id'],
                                  'generated_output': values[model['source_column_index']],
                                  'original_row_number': row_number,
                                  'source_column_index': model['source_column_index'], 'source_column_name': model['source_column_name']})
        return {'mapping': configuration.model_dump(), 'models': models, 'cases': cases, 'responses': responses}
    finally:
        if text: text.detach()
        if workbook: workbook.close()
