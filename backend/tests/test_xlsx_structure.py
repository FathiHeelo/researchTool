from datetime import date
from io import BytesIO
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app

client = TestClient(app)


def upload(rows):
    workbook = Workbook()
    workbook.active.title = "Other"
    for row in [["ignored", "columns", "here"], [1, 2, 3], [4, 5, 6]]:
        workbook.active.append(row)
    selected = workbook.create_sheet("Selected")
    for row in rows:
        selected.append(row)
    # Formatting-only cells must not inflate the tabular structure.
    selected["Z100"].number_format = "0.00"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    response = client.post("/datasets/upload", files={"file": ("data.xlsx", stream.getvalue())}, data={"sheet_name": "Selected"})
    assert response.status_code == 200, response.text
    data = response.json()
    assert set(data) == {"filename", "extension", "size", "sheet_names", "sheet_count", "selected_sheet", "total_row_count", "column_count", "columns", "inferred_data_types", "preview_rows"}
    return data


def test_normal_headers_order_and_rows():
    data = upload([[" Z ", "Name", "A"], [1, "study", True], [2, None, False]])
    assert data["columns"] == [" Z ", "Name", "A"]
    assert data["column_count"] == 3
    assert data["total_row_count"] == 2


def test_header_only():
    data = upload([["a", "b"]])
    assert data["columns"] == ["a", "b"]
    assert data["column_count"] == 2
    assert data["total_row_count"] == 0


@pytest.mark.parametrize("rows", [[], [[None]], [[" "]]])
def test_empty_selected_sheet(rows):
    data = upload(rows)
    assert data["columns"] == []
    assert data["column_count"] == data["total_row_count"] == 0


def test_unicode_and_duplicate_headers():
    data = upload([["بيانات", "Résumé", "بيانات"], [1, 2, 3]])
    assert data["columns"] == ["بيانات", "Résumé", "بيانات"]
    assert data["column_count"] == 3


def test_blank_headers_and_wider_data():
    data = upload([["a", None, "c"], [1, 2, None, 4]])
    assert data["columns"] == ["a", None, "c", None]
    assert data["column_count"] == 4
    assert data["total_row_count"] == 1


def test_null_rows_and_missing_cells():
    data = upload([["a", "b"], [1], [None, None], [None, 2], [0, False]])
    assert data["columns"] == ["a", "b"]
    assert data["total_row_count"] == 3


def test_blank_first_row_remains_header():
    data = upload([[None, None], [1, 2]])
    assert data["columns"] == [None, None]
    assert data["total_row_count"] == 1


def test_preamble_rows_before_real_header_are_supported():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Prompt Engineering Dataset"
    worksheet.append(["Few-shot Prompts - 130 Records"])
    worksheet.append(["Description"])
    worksheet.append(["Records: 130"])
    worksheet.append([None, None, None, None, "Actual Output"])
    worksheet.append(["ID", "Original User Prompt", "Generated Expectations", "Few-shot Prompt", "Model Alpha"])
    worksheet.append(["01", "Requirement", 'expect_column_to_exist(column="a")', "Prompt", 'expect_column_to_exist(column="a")'])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    response = client.post("/datasets/upload", files={"file": ("data.xlsx", stream.getvalue())}, data={"sheet_name": "Prompt Engineering Dataset"})
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["columns"][:5] == ["ID", "Original User Prompt", "Generated Expectations", "Few-shot Prompt", "Model Alpha"]
    assert data["total_row_count"] == 1
    assert data["preview_rows"][0]["row_index"] == 1

    mapping = {"assignments": [
        {"index": 0, "role": "case_id"},
        {"index": 1, "role": "requirement"},
        {"index": 2, "role": "expected_rule"},
        {"index": 3, "role": "metadata", "label": "Prompt Strategy"},
        {"index": 4, "role": "model"},
    ]}
    import_response = client.post(
        "/datasets/import",
        files={"file": ("data.xlsx", stream.getvalue())},
        data={"sheet_name": "Prompt Engineering Dataset", "mapping": json.dumps(mapping)},
    )
    assert import_response.status_code == 200, import_response.text
    imported = import_response.json()
    assert imported["cases"][0]["case_id"] == "01"
    assert imported["cases"][0]["original_row_number"] == 6
    assert imported["responses"][0]["generated_output"] == 'expect_column_to_exist(column="a")'


def test_nontext_headers_are_json_safe():
    data = upload([[42, True, date(2026, 9, 7)], [1, 2, 3]])
    assert data["columns"] == [42, True, "2026-09-07T00:00:00"]
