import csv
import io
import json
import pytest
from openpyxl import load_workbook
from test_runs import client, benchmark
from app.evaluation.engine import evaluate_pair
from app.evaluation.metrics import defaults, overall
from app.evaluation.ground_truth import validate_ground_truth


BASE = '/projects/1/runs'


def test_configuration_snapshot_and_accepted_scores(client):
    profile = client.get(BASE + '/settings/metrics').json()
    assert len(profile['metrics']) == 4
    assert overall({'syntax': 1, 'execution': 1, 'semantic': .5, 'completeness': 1}, defaults()) == 87.5
    client.post(BASE, json=benchmark())
    for m in profile['metrics']:
        m['enabled'] = m['key'] == 'semantic'
    assert client.post(BASE + '/settings/metrics', json=profile).status_code == 200
    client.post(BASE, json=benchmark())
    first = client.post(BASE + '/1/results/1/review', json={'override_scores': {'semantic': .5}, 'reason': 'Verified'}).json()
    second = client.post(BASE + '/2/results/5/review', json={'override_scores': {'semantic': .5}, 'reason': 'Verified'}).json()
    assert first['accepted_overall_accuracy'] == 87.5
    assert second['accepted_overall_accuracy'] == 50
    assert first['automated_scores']['semantic'] == 1
    assert client.post(BASE + '/settings/metrics/reset').status_code == 200
    assert client.get(BASE + '/2/results/5/review').json()['accepted_overall_accuracy'] == 50


@pytest.mark.parametrize('change', ['zero', 'duplicate', 'scale', 'negative', 'nonfinite'])
def test_invalid_metric_configuration(client, change):
    metrics = defaults()
    if change == 'zero':
        for m in metrics: m['weight'] = 0
    if change == 'duplicate': metrics.append(metrics[0])
    if change == 'scale': metrics[0]['max_score'] = 0
    if change == 'negative': metrics[0]['weight'] = -1
    if change == 'nonfinite': metrics[0]['weight'] = 'NaN'
    assert client.post(BASE + '/settings/metrics', json={'metrics': metrics}).status_code == 422


def test_arbitrary_metric_scale(client, monkeypatch):
    from app.evaluation import batch
    original = batch.evaluate_pair
    def evaluator(a, b):
        result = original(a,b); result['component_scores']['custom-research'] = 5; return result
    monkeypatch.setattr(batch, 'evaluate_pair', evaluator)
    config = {'key': 'custom-research', 'name': 'Custom Research', 'min_score': 0, 'max_score': 10, 'weight': 7}
    assert client.post(BASE + '/settings/metrics', json={'metrics': [config]}).status_code == 200
    client.post(BASE, json=benchmark())
    row = client.post(BASE + '/1/results/1/review', json={'override_scores': {'custom-research': 9}, 'reason': 'Scale validated'}).json()
    assert row['accepted_overall_accuracy'] == 90
    assert client.post(BASE + '/1/results/1/review', json={'override_scores': {'custom-research': 11}, 'reason': 'bad'}).status_code == 400


def research_benchmark():
    data = benchmark()
    data['cases'][0]['case_id'] = data['cases'][1]['case_id'] = 'matched'
    for i, case in enumerate(data['cases']):
        case['metadata'] = [{'label': 'Experiment', 'value': ['Alpha', 'Beta'][i]}, {'label': 'DimensionX', 'value': 'Novel'}, {'label': 'عربي', 'value': 'بحث'}]
    return data


def test_research_analytics_and_scopes(client):
    client.post(BASE, json=research_benchmark())
    client.post(BASE + '/1/results/1/review', json={'override_scores': {'semantic': 0}, 'reason': 'Checked', 'error_tags': ['A', 'B', 'A']})
    data = client.get(BASE + '/dashboard/research?strategy_key=Experiment&dimension_key=DimensionX&group_key=عربي').json()
    assert data['response_count'] == 4
    assert data['reliability'] == 50
    assert data['hallucination_rate'] == 50
    assert data['hallucinated_functions'] == {'expect_invented': 2}
    delta = next(d for d in data['strategy_deltas'] if d['model'] == 'Future Model')
    assert delta['matched_case_count'] == 1
    assert delta['delta'] == 25
    assert len(data['dimensions']) == len(data['groups']) == 2
    assert len(data['dimension_summary']) == 1
    assert data['agreement'][0]['comparable_count'] == 1
    errors = {e['tag']: e for e in data['errors']}
    assert errors['A']['count'] == 1 and errors['A']['percentage'] == 25
    assert client.get(BASE + '/dashboard/research?run_id=1&model=Another').json()['response_count'] == 2


def test_unparseable_agreement_and_missing_metadata(client):
    data = benchmark()
    for response in data['responses']: response['generated_output'] = '!!!'
    client.post(BASE, json=data)
    result = client.get(BASE + '/dashboard/research').json()
    assert result['agreement'][0]['agreement_rate'] is None
    assert result['agreement'][0]['non_comparable_count'] == 2
    assert result['strategies'] == result['dimensions'] == []


def test_empty_analytics_and_ground_truth_filter(client):
    assert client.get(BASE + '/dashboard/research').json()['response_count'] == 0
    data = benchmark(); data['cases'][0]['expected_rule'] = 'expect_invented()'
    client.post(BASE, json=data)
    result = client.get(BASE + '/dashboard/research?include_ground_truth_warnings=false').json()
    assert result['response_count'] == result['filtered_results'] == 2
    assert result['filtered_cases'] == 1
    summary = client.get(BASE + '/dashboard/summary?include_ground_truth_warnings=false').json()
    assert summary['total_evaluated_responses'] == 2


def test_import_to_export_integration(client):
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(['Case', 'Requirement', 'Expected', 'Model Ω', 'Experiment'])
    writer.writerow(['C1', 'بحث', 'expect_column_to_exist("a")', 'expect_column_to_exist("a")', 'Custom'])
    writer.writerow(['C2', '=not_a_formula', 'expect_invented()', 'expect_column_to_exist("a")', 'Custom'])
    roles = ['case_id', 'requirement', 'expected_rule', 'model', 'metadata']
    mapping = {'assignments': [{'index': i, 'role': role} for i, role in enumerate(roles)]}
    imported = client.post('/datasets/import', files={'file': ('research.csv', stream.getvalue().encode())}, data={'mapping': json.dumps(mapping)})
    assert imported.status_code == 200
    run = client.post(BASE, json=imported.json())
    assert run.status_code == 201
    rows = client.get(BASE + '/1/results').json()['items']
    assert rows[1]['ground_truth_warning']
    reviewed = client.post(BASE + '/1/results/1/review', json={'override_scores': {'semantic': .5}, 'reason': 'Research reason', 'actor': 'Researcher Ω'}).json()
    assert reviewed['automated_scores']['semantic'] == 1
    assert reviewed['accepted_overall_accuracy'] == 87.5
    assert client.get(BASE + '/dashboard/research').json()['response_count'] == 2
    response = client.get(BASE + '/1/export?strategy_key=Experiment')
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    assert 'Manual Overrides' in book.sheetnames and 'Metric Configuration' in book.sheetnames
    assert 'DQ Dimension Analysis' not in book.sheetnames
    ws = book['Detailed Results']; headers = [c.value for c in ws[1]]
    assert ws.cell(2, headers.index('Source Row')+1).value == 2
    assert ws.cell(2, headers.index('Accepted Overall Accuracy')+1).value == 87.5
    assert ws.cell(3, headers.index('Requirement')+1).data_type == 's'
    assert ws.freeze_panes == 'A2'
    assert book['Manual Overrides'].max_row > 1
    csv_response = client.get(BASE + '/1/export?format=csv')
    assert 'بحث' in csv_response.content.decode('utf-8-sig')
    assert "'=not_a_formula" in csv_response.content.decode('utf-8-sig')
    filtered = load_workbook(io.BytesIO(client.get(BASE + '/1/export?include_ground_truth_warnings=false').content))
    assert filtered['Detailed Results'].max_row == 2
    assert client.get(BASE + '/1/export?format=invalid').status_code == 400


# Deterministic benchmark: semantics and GX execution are separate assertions.
REGRESSION = [
    ('exact', 'expect_column_to_exist("a")', 'expect_column_to_exist("a")', 1, None),
    ('order', 'expect_column_to_exist("a")\nexpect_column_to_exist("b")', 'expect_column_to_exist("b")\nexpect_column_to_exist("a")', 1, None),
    ('text alias', 'expect_column_values_to_be_of_type("a", "text")', 'expect_column_values_to_be_of_type("a", "str")', 1, None),
    ('integer alias', 'expect_column_values_to_be_of_type("a", "integer")', 'expect_column_values_to_be_of_type("a", "int")', 1, None),
    ('fraction type', 'expect_column_values_to_be_of_type("a", "float")', 'expect_column_values_to_be_of_type("a", "decimal")', .5, 'WRONG_TYPE'),
    ('boundary', 'expect_column_values_to_be_between("a", min_value=0)', 'expect_column_values_to_be_between("a", min_value=0, strict_min=True)', .5, 'WRONG_BOUNDARY'),
    ('missing', 'expect_column_to_exist("a")\nexpect_column_to_exist("b")', 'expect_column_to_exist("a")', .5, 'MISSING_CONDITION'),
    ('invented', 'expect_column_to_exist("a")', 'expect_invented("a")', 0, 'UNKNOWN_FUNCTION'),
    ('regex subset', 'expect_column_values_to_match_regex("a", "^(a|b)$")', 'expect_column_values_to_match_regex("a", "^a$")', .5, 'REGEX_STRICTER'),
    ('extra', 'expect_column_to_exist("a")', 'expect_column_to_exist("a")\nexpect_column_to_exist("b")', .5, 'EXTRA_CONSTRAINT'),
]


@pytest.mark.parametrize('name,expected,generated,semantic,error', REGRESSION, ids=[r[0] for r in REGRESSION])
def test_research_regression(name, expected, generated, semantic, error):
    result = evaluate_pair(expected, generated)
    assert result['semantic'] == semantic
    if error: assert error in {e['code'] for e in result['errors']}
    if name == 'invented': assert result['syntax'] == 1 and result['execution'] == 0 and result['hallucination_detected']
    if name == 'missing': assert result['completeness'] < 1


def test_parser_recovery_and_reference_warning():
    assert evaluate_pair('expect_column_to_exist("a")', '!!!\nexpect_column_to_exist("a")')['syntax'] == 0
    assert validate_ground_truth('expect_invented()')[0]['warning_type'] == 'UNKNOWN_EXPECTATION'
    assert not validate_ground_truth('expect_column_to_exist("a")\nexpect_column_to_exist("b")')
    assert validate_ground_truth('expect_column_to_exist("a")\nexpect_column_to_exist("a")')[0]['warning_type'] == 'SUSPICIOUS_CONSTRAINT'
