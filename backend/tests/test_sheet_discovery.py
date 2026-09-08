from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app

client = TestClient(app)


def upload_workbook(workbook):
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    content = stream.getvalue()
    response = client.post("/datasets/upload", files={"file": ("study.xlsx", content)})
    return response, len(content)


@pytest.mark.parametrize("names", [
    ["Sheet1"],
    ["Zebra", "Alpha", "Middle"],
    [" Research Data ", "Notes and references"],
    ["بيانات", "研究", "Résultats"],
])
def test_sheet_names_and_order(names):
    workbook = Workbook()
    workbook.active.title = names[0]
    workbook.active.append(["header"])
    for name in names[1:]:
        workbook.create_sheet(name)
    response, size = upload_workbook(workbook)
    assert response.status_code == 200
    assert response.json() == {
        "filename": "study.xlsx", "extension": ".xlsx", "size": size,
        "sheet_names": names, "sheet_count": len(names),
    }


@pytest.mark.parametrize("state", ["hidden", "veryHidden"])
def test_hidden_and_empty_worksheets_included(state):
    workbook = Workbook()
    workbook.active.title = "Empty visible"
    workbook.create_sheet("Empty hidden").sheet_state = state
    workbook.create_sheet("Data").append(["header"])
    response, _ = upload_workbook(workbook)
    assert response.status_code == 200
    assert response.json()["sheet_names"] == ["Empty visible", "Empty hidden", "Data"]
    assert response.json()["sheet_count"] == 3


def test_entirely_empty_workbook_still_rejected():
    response, _ = upload_workbook(Workbook())
    assert response.status_code == 400
    assert response.json()["code"] == "unreadable_dataset"


def test_csv_response_unchanged():
    content = b"name,value\nStudy,1\n"
    response = client.post("/datasets/upload", files={"file": ("data.csv", content)})
    assert response.status_code == 200
    assert response.json() == {
        "filename": "data.csv", "extension": ".csv", "size": len(content),
        "total_row_count": 1, "column_count": 2, "columns": ["name", "value"],
        "inferred_data_types": {"name": "text", "value": "integer"},
        "preview_rows": [{"name": "Study", "value": "1"}],
    }
