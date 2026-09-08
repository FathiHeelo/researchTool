from io import BytesIO
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app

client = TestClient(app)


def workbook_bytes(rows=()):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def assert_invalid(filename, content, code):
    response = client.post("/datasets/upload", files={"file": (filename, content)})
    assert response.status_code == 400
    assert response.json()["code"] == code
    assert isinstance(response.json()["detail"], str)
    assert response.json()["detail"]
    assert "Traceback" not in response.text


@pytest.mark.parametrize("content", [b'a,b\n"unclosed,b', b'a,b\n1,2,3\n', b'a,b\n"a"x,b', b'a,b\n\xff,2', b'a,b\n\x00,2'])
def test_malformed_csv(content):
    assert_invalid("data.csv", content, "malformed_csv")


@pytest.mark.parametrize("content", [b" \n\t\n", b",,\n,,\n", b"\xef\xbb\xbf"])
def test_csv_without_usable_data(content):
    assert_invalid("data.csv", content, "unreadable_dataset")


@pytest.mark.parametrize("content", [
    '\ufeffname,note\r\nStudy,café\r\n'.encode('utf-8'),
    b'name,note\nStudy,"quoted, value"\n',
    b'name,note\nStudy,"multiple\nlines"\n',
    b'name,note\n\nStudy,\n',
])
def test_common_csv_formats(content):
    response = client.post("/datasets/upload", files={"file": ("data.csv", content)})
    assert response.status_code == 200
    assert {key: response.json()[key] for key in ("filename", "extension", "size")} == {"filename": "data.csv", "extension": ".csv", "size": len(content)}


def test_fake_xlsx():
    assert_invalid("data.xlsx", b"not an Excel workbook", "malformed_xlsx")


def test_truncated_xlsx():
    assert_invalid("data.xlsx", workbook_bytes([["name"], ["Study"]])[:100], "malformed_xlsx")


def test_malformed_worksheet_xml():
    source = BytesIO(workbook_bytes([["name"]]))
    target = BytesIO()
    with ZipFile(source) as original, ZipFile(target, "w") as changed:
        for item in original.infolist():
            changed.writestr(item.filename, b"<broken" if item.filename == "xl/worksheets/sheet1.xml" else original.read(item.filename))
    assert_invalid("data.xlsx", target.getvalue(), "malformed_xlsx")


@pytest.mark.parametrize("rows", [[], [[" ", "\t"]]])
def test_empty_workbook(rows):
    assert_invalid("data.xlsx", workbook_bytes(rows), "unreadable_dataset")


def test_workbook_with_data_on_second_sheet():
    workbook = Workbook()
    workbook.create_sheet("Data").append([0, False])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    response = client.post("/datasets/upload", files={"file": ("data.xlsx", stream.getvalue())})
    assert response.status_code == 200
    assert set(response.json()) == {"filename", "extension", "size", "sheet_names", "sheet_count"}
    assert response.json()["sheet_names"] == ["Sheet", "Data"]


@pytest.mark.parametrize("filename,content,code", [("data.csv", b"", "empty_file"), ("data.xlsx", b"", "empty_file"), ("data.txt", b"abc", "unsupported_file")])
def test_structured_validation_errors(filename, content, code):
    assert_invalid(filename, content, code)
